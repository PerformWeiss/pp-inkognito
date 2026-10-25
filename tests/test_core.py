"""Kurztest mit erfundenen Daten: python -m tests.test_core"""
import io
import os
import re
import tempfile

os.environ["ANONYMISIERUNGSTOOL_HOME"] = tempfile.mkdtemp()
os.environ["ANONYMISIERUNGSTOOL_AUSGABE"] = tempfile.mkdtemp()

import docx  # noqa: E402

from anonymisierungstool.detect import analyze  # noqa: E402
from anonymisierungstool.documents import DocxDoc, PdfDoc, open_document  # noqa: E402
from anonymisierungstool.keyfile import dump_key, load_key  # noqa: E402
from anonymisierungstool.pseudonymize import FAKE_FIRST, Item, build_mapping  # noqa: E402
from anonymisierungstool.textutil import term_regex  # noqa: E402


def sample_docx():
    d = docx.Document()
    d.core_properties.author = "Anna Beispiel"
    d.sections[0].header.paragraphs[0].text = "Anna Beispiel · Lebenslauf"
    d.add_paragraph("Lebenslauf")
    p = d.add_paragraph()
    p.add_run("Anna ").bold = True          # Name über zwei Runs verteilt
    p.add_run("Beispiel").bold = True
    d.add_paragraph("Musterstraße 12, 80331 München")
    d.add_paragraph("Tel. +49 151 23456789 · anna.beispiel@example.org · www.linkedin.com/in/annabeispiel")
    d.add_paragraph("Geboren am 14.03.1988 in Augsburg")
    d.add_paragraph("03/2019 – 08/2021  Projektleiterin bei Kowalski Bau GmbH, Warschau")
    d.add_paragraph("2015 - 2018  Müller & Söhne KG, Stuttgart")
    d.add_paragraph("Frau Beispiel verfügt über Erfahrung im Projektmanagement.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Geburtsdatum"
    t.rows[0].cells[1].text = "14.03.1988"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def sample_pdf():
    import fitz

    d = fitz.open()
    pg = d.new_page()
    pg.draw_rect(fitz.Rect(0, 0, 160, 842), color=None, fill=(0.85, 0.9, 0.97))  # farbige Seitenleiste
    pg.insert_text((20, 60), "Anna Beispiel", fontname="hebo", fontsize=14)
    pg.insert_text((20, 80), "anna.beispiel@example.org", fontname="helv", fontsize=9)
    pg.insert_text((20, 95), "+49 151 23456789", fontname="helv", fontsize=9)
    pg.insert_text((190, 60), "Lebenslauf", fontname="hebo", fontsize=16)
    pg.insert_text((190, 90), "Geboren am 14.03.1988 in Augsburg", fontname="helv", fontsize=10)
    pg.insert_text((190, 110), "Musterstrasse 12, 80331 Muenchen", fontname="helv", fontsize=10)
    pg.insert_text((190, 140), "03/2019 - 08/2021", fontname="helv", fontsize=10)
    pg.insert_text((300, 140), "Projektleiterin, Kowalski Bau GmbH, Warschau", fontname="helv", fontsize=10)
    pg.draw_line((190, 150), (560, 150))
    pg.insert_text((190, 170), "Frau Beispiel hat Erfahrung. Webentwicklung und Annahme bleiben.", fontname="helv", fontsize=10)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 50), False)
    pix.set_rect(pix.irect, (200, 50, 50))
    pg.insert_image(fitz.Rect(20, 120, 100, 220), pixmap=pix)
    d.set_metadata({"author": "Anna Beispiel", "title": "CV Anna Beispiel"})
    return d.tobytes()


def test_pdf(_unused):
    import fitz

    data = sample_pdf()
    doc = open_document("cv.pdf", data)
    findings, _ = analyze(doc.text())
    for f in findings:
        print("  PDF-Fund: %-14s %s" % (f["category"], f["text"]))
    items = [Item(f["text"], f["category"]) for f in findings if f["checked"]] + [Item("Kowalski Bau GmbH", "firma")]
    mapping, _ = build_mapping(items, "verschieben")
    out, ext, _ = doc.export(term_regex(mapping), mapping, True)
    left = doc.leftover_check(out, ext, list(mapping))
    res = fitz.open(stream=out, filetype="pdf")
    txt = res[0].get_text()
    print("--- PDF-Ausgabe ---\n" + txt)
    print("Metadaten:", res.metadata.get("author"), "| Bilder:", len(res[0].get_images()), "| Übrig:", left)
    assert not left and "Beispiel" not in txt and "anna." not in txt
    assert "Webentwicklung" in txt and "Annahme" in txt, "Wortteile wurden fälschlich ersetzt"
    assert "Foto entfernt" in txt, "Platzhalter für das entfernte Foto fehlt"
    pix = res[0].get_pixmap(clip=fitz.Rect(30, 150, 90, 200), dpi=72)
    assert pix.pixel(10, 10)[0] < 190 or pix.pixel(10, 10)[1] > 40, "Foto ist noch sichtbar"
    assert not res.metadata.get("author")


def test_scan_pages():
    import fitz

    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 300), False)
    pix.set_rect(pix.irect, (240, 240, 240))
    d = fitz.open()
    d.new_page().insert_text((50, 80), "Textseite mit Inhalt", fontname="helv", fontsize=12)
    d.new_page().insert_image(fitz.Rect(0, 0, 400, 600), pixmap=pix)   # Scan-Seite
    d.new_page()                                                        # leere Seite, kein Scan
    mixed = open_document("gemischt.pdf", d.tobytes())
    assert mixed.scan_pages == [2], mixed.scan_pages
    assert any("Seite 2 dieser PDF ist ein Scan" in w for w in mixed.warnings), mixed.warnings

    s = fitz.open()
    s.new_page().insert_image(fitz.Rect(0, 0, 400, 600), pixmap=pix)
    scan = open_document("scan.pdf", s.tobytes())
    assert scan.scan_pages == [1] and any("Alle Seiten" in w for w in scan.warnings), scan.warnings

    plain = open_document("text.pdf", sample_pdf())
    assert plain.scan_pages == [] and not any("Scan" in w for w in plain.warnings)

    from anonymisierungstool import server
    client = server.app.test_client()
    client.set_cookie("tok", server.TOKEN, domain="localhost")
    r = client.post("/api/upload", data={"file": (io.BytesIO(b"Dies ist ein Text ohne jede Angabe zu einer Person."), "a.txt")},
                    headers={"Host": "localhost"}).get_json()
    assert any("kein Name gefunden" in w for w in r["warnings"]), r["warnings"]
    print("Scan- und Namenshinweise geprüft.")


def test_titel():
    """Zwei Dokumente unter einem Titel: gleiche Pseudonyme, ergänzter Schlüssel, Rückspielen über den Titel."""
    from anonymisierungstool import server, titel

    assert titel.folder_name("../../etc/passwd") == "etc-passwd" and titel.folder_name("Bewerbung: Anna?") == "Bewerbung- Anna"
    try:
        titel.folder_name("  ..  ")
        raise AssertionError("leerer Titel akzeptiert")
    except titel.InvalidTitle:
        pass

    # Pseudonyme über Dokumente hinweg
    key = {"entries": []}
    m1, _ = build_mapping([Item("Anna Beispiel", "name"), Item("Kowalski Bau GmbH", "firma"), Item("14.03.1988", "geburtsdatum")], "verschieben", key=key)
    m2, _ = build_mapping([Item("Anna Beispiel", "name"), Item("Beispiel", "name"), Item("Bernd Muster", "name"), Item("Müller KG", "firma"),
                           Item("20.05.1990", "geburtsdatum")], "verschieben", key=key)
    anna = m1["Anna Beispiel"].split()
    assert m2["Anna Beispiel"] == m1["Anna Beispiel"] and len(anna) == 2 and anna[1].startswith("Muster")
    assert m2["Beispiel"] == anna[1], "Nachname allein muss denselben erfundenen Nachnamen bekommen"
    bernd = m2["Bernd Muster"].split()
    assert bernd[0] != anna[0] and bernd[1] != anna[1] and m2["Müller KG"] == "Musterfirma B"
    assert m2["20.05.1990"][-4:] == str(int("1990") + (int(m1["14.03.1988"][-4:]) - 1988)), "Datumsverschiebung nicht einheitlich"
    assert len(key["entries"]) == 7 and len({e["original"] for e in key["entries"]}) == 7

    # Schwärzen in einem späteren Dokument ändert das Pseudonym im Schlüssel nicht
    m3, _ = build_mapping([Item("Anna Beispiel", "name", "schwaerzen")], "verschieben", key=key)
    assert m3["Anna Beispiel"] == "[geschwärzt]" and {e["original"]: e for e in key["entries"]}["Anna Beispiel"]["replacement"] == m1["Anna Beispiel"]

    # Ende zu Ende über den Server
    client = server.app.test_client()
    client.set_cookie("tok", server.TOKEN, domain="localhost")
    hdr = {"Host": "localhost"}
    title = "Test Anna"
    for text in (b"Anna Beispiel wohnt in Berlin. Mail: anna.beispiel@example.org", b"Bewerbung von Anna Beispiel und Bernd Muster."):
        up = client.post("/api/upload", data={"file": (io.BytesIO(text), "Lebenslauf_Anna_Beispiel.txt")}, headers=hdr).get_json()
        items = [{"text": t, "category": "name"} for t in ("Anna Beispiel", "Bernd Muster") if t.encode() in text]
        ex = client.post("/api/export", json={"doc_id": up["doc_id"], "title": title, "items": items}, headers=hdr).get_json()
        assert "error" not in ex, ex
        assert "Anna" not in ex["file"] and "Beispiel" not in ex["file"] and ex["file"].startswith("Lebenslauf_") and "_anonymisiert" in ex["file"], ex["file"]
    folder = titel.folder(title)
    names = sorted(p.name for p in folder.iterdir())
    anon = [n for n in names if n.endswith(".txt")]
    assert len(anon) == 2 and "SCHLUESSEL_nicht_in_die_KI_laden.anonkey" in names, names
    assert not any("Anna" in n or "Beispiel" in n for n in names), names
    assert all(titel.file_kind(n) == "anonymisiert" for n in anon)
    out2 = (folder / ex["file"]).read_text(encoding="utf-8")
    assert "Beispiel" not in out2 and "Anna" not in out2 and "Bernd" not in out2, out2
    assert title in titel.list_titles(with_key=True)
    back = client.post("/api/restore", data={"file": (io.BytesIO(out2.encode()), "ki.txt"), "title": title}, headers=hdr).get_json()
    assert "error" not in back and not back["unknown"], back
    assert (folder / back["file"]).read_text(encoding="utf-8") == "Bewerbung von Anna Beispiel und Bernd Muster."
    assert back["file"].endswith("_wiederhergestellt.txt") and titel.file_kind(back["file"]) == "wiederhergestellt", back["file"]
    from anonymisierungstool.pseudonymize import anonymize_filename
    fm = {"Max Mustermann": "Horst Muster", "Mustermann": "Muster"}
    assert anonymize_filename("CV_Max_Mustermann_2024", fm) == "CV_Horst_Muster_2024"
    assert anonymize_filename("LebenslaufMaxMustermann", fm) == "LebenslaufHorst_Muster" and anonymize_filename("cv_003", fm) == "cv_003"
    assert "/" not in anonymize_filename("a/b:c", {}) and anonymize_filename("..", {}) == "Dokument"
    bad = client.post("/api/export", json={"doc_id": up["doc_id"], "title": "  ", "items": items}, headers=hdr)
    assert bad.status_code == 400
    print("Titel, gemeinsamer Schlüssel und Rückspielen geprüft.")


def test_teil_a():
    """Datumsmodus je Eintrag, Kategorie Projekt, Bilder zählen, Scan-Seiten nicht leeren."""
    import fitz

    m, _ = build_mapping([Item("14.03.1988", "geburtsdatum"), Item("01.02.1990", "geburtsdatum", mode="schwaerzen"),
                          Item("Projekt Atlas", "projekt")])
    assert m["14.03.1988"] == "1988" and m["01.02.1990"] == "[geschwärzt]" and m["Projekt Atlas"] == "Projekt A", m
    m, _ = build_mapping([Item("14.03.1988", "geburtsdatum", mode="verschieben")])
    assert m["14.03.1988"] != "1988" and m["14.03.1988"].startswith("14.03."), m

    pdf = open_document("cv.pdf", sample_pdf())
    assert pdf.image_count == 1, pdf.image_count
    d = DocxDoc(sample_docx())
    assert d.image_count == 0

    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 300), False)
    pix.set_rect(pix.irect, (120, 120, 120))
    src = fitz.open()
    src.new_page().insert_text((50, 80), "Textseite mit Anna Beispiel", fontname="helv", fontsize=12)
    src.new_page().insert_image(fitz.Rect(0, 0, 400, 600), pixmap=pix)   # Scan-Seite
    doc = open_document("gemischt.pdf", src.tobytes())
    assert doc.image_count == 0 and doc.scan_pages == [2]
    mapping = {"Anna Beispiel": "Person A"}
    out, _, _ = doc.export(term_regex(mapping), mapping, True)
    res = fitz.open(stream=out, filetype="pdf")
    assert len(res[1].get_images()) == 1, "Scan-Bild wurde entfernt"
    assert res[1].get_pixmap(dpi=20).samples.count(120) > 100, "Scan-Seite ist leer"
    print("Teil A (Datumsmodus, Projekt, Bilderzahl, Scan-Seiten) geprüft.")


def test_teil_c():
    """Erkennung für online erstellte Lebensläufe: ISO-Datum, Straße ohne Endung, Versalien, Satzanfänge, Reihenfolge."""
    import fitz
    from anonymisierungstool.detect import analyze_all
    from anonymisierungstool.documents import _page_text

    text = ("Max Mustermann\nWiesenblick 5a-12345 Neustadt am Rübenberge | 0151 2345678\nmax.mustermann@example.org\n1980-05-17\n"
            "Berufserfahrung\nMÜNCHEN\nMUSTERFIRMA GMBH\n• Leitete 50 Mitarbeiter und verbesserte die Teamleistung.\n"
            "• Konzipiere Seminare zur Kompetenzentwicklung.")
    found, _ = analyze_all(text)
    by = {(f["category"], f["text"]): f for f in found}
    assert by[("geburtsdatum", "1980-05-17")]["checked"], "ISO-Geburtsdatum nicht erkannt"
    assert ("adresse", "Wiesenblick 5a") in by and ("adresse", "12345 Neustadt am Rübenberge") in by, [k for k in by if k[0] == "adresse"]
    assert ("ort", "MÜNCHEN") in by and ("firma", "MUSTERFIRMA GMBH") in by
    assert ("ort", "NEUSTADT AM RÜBENBERGE") not in by and not [k for k in by if k[1] in ("Leitete", "Konzipiere")], list(by)
    assert [f["text"] for f in found if f["category"] == "name"][0] == "Max Mustermann", "Reihenfolge der Namen"
    m, _e = build_mapping([Item("1980-05-17", "geburtsdatum")])
    assert m["1980-05-17"] == "1980"
    m, _e = build_mapping([Item("1980-05-17", "geburtsdatum", mode="verschieben")])
    assert m["1980-05-17"].endswith("-05-17") and m["1980-05-17"][:4] != "1980"

    d = fitz.open()
    pg = d.new_page()
    pg.insert_text((50, 200), "Zeile unten", fontname="helv", fontsize=11)        # zuerst gezeichnet, steht aber weiter unten
    pg.insert_text((50, 100), "Spalte links", fontname="helv", fontsize=11)
    pg.insert_text((300, 100), "Spalte rechts", fontname="helv", fontsize=11)
    lines = _page_text(d[0]).split("\n")
    assert lines == ["Spalte links\tSpalte rechts", "Zeile unten"], lines
    print("Teil C (Erkennung, Lesereihenfolge) geprüft.")


def test_teil_b():
    """Bewerber/in, Muster-Adressen, Überschreiben und gleiche Vorschau wie Export."""
    from anonymisierungstool import server
    from anonymisierungstool.detect import analyze_all, guess_gender

    assert guess_gender("Fabian Beispiel", "") == "m" and guess_gender("Anna Beispiel", "") == "f" and guess_gender("Xyz Beispiel", "") == ""
    assert guess_gender("Xyz Beispiel", "Sehr geehrte Frau Beispiel,") == "f"
    text = "Anna Beispiel\nMusterweg 3, 80331 München\nanna.beispiel@example.org\nFrau Beispiel arbeitet bei Kowalski Bau GmbH."
    found, _ = analyze_all(text)
    main = [f for f in found if f.get("main")]
    assert len(main) == 1 and main[0]["text"] == "Anna Beispiel" and main[0]["gender"] == "f", [(f["text"], f.get("main")) for f in found]

    key = {"entries": []}
    items = [Item(f["text"], f["category"], main=f["main"], gender=f["gender"]) for f in found if f["checked"]]
    m, _e = build_mapping(items, key=key)
    erika = m["Anna Beispiel"].split()
    assert erika[0] in FAKE_FIRST["f"] and erika[1].startswith("Muster") and m.get("Beispiel", erika[1]) == erika[1]
    assert m["Musterweg 3"] if "Musterweg 3" in m else True
    assert m["anna.beispiel@example.org"] == "muster@example.com"
    adr = [v for k, v in m.items() if "München" in k]
    assert adr and adr[0].startswith("12345 "), m
    m2, _e = build_mapping([Item("Bernd Muster", "name", main=True)], key=key)       # zweite Hauptperson im selben Titel: kein zweiter Bewerber
    assert m2["Bernd Muster"].split()[0] != erika[0] and m2["Bernd Muster"].split()[1] != erika[1]
    m3, _e = build_mapping([Item("Bernd Muster", "name", override="Kollege X")], key=key)
    from anonymisierungstool.restore import build_reverse
    k2 = {"entries": []}
    build_mapping([Item("Hans Beispiel", "name"), Item("Beispiel", "name"), Item("Hans", "name")], key=k2)
    rev, lost = build_reverse(k2)
    assert sorted(rev.values()) == ["Beispiel", "Hans", "Hans Beispiel"] and lost == 0, (rev, lost)
    fritz, _e = build_mapping([Item("Fritz Müller", "name", main=True, gender="m"), Item("Fritz", "name"), Item("Müller", "name"), Item("MÜLLER", "name")])
    f1, f2 = fritz["Fritz Müller"].split()
    assert fritz["Fritz"] == f1 and fritz["Müller"] == f2 and fritz["MÜLLER"] == f2.upper() and f1 in FAKE_FIRST["m"], fritz
    blocked, _e = build_mapping([Item("Fritz", "name", gender="m")], avoid={n.lower() for n in FAKE_FIRST["m"] if n != "Horst"})
    assert blocked["Fritz"] == "Horst", "Vornamen aus dem Dokument dürfen nicht als Pseudonym dienen"
    own, _e = build_mapping([Item("Anna Schmidt", "name")])
    assert own["Anna Schmidt"].split()[0] != "Anna", "Pseudonym darf nicht dem Original gleichen"
    m4, _e = build_mapping([Item("Hans Beispiel", "name", override="Kandidat")])
    assert m4["Hans Beispiel"] == "Kandidat", m4
    assert m3["Bernd Muster"] == "Kollege X" and {e["original"]: e for e in key["entries"]}["Bernd Muster"]["replacement"] == "Kollege X"

    client = server.app.test_client()
    client.set_cookie("tok", server.TOKEN, domain="localhost")
    hdr = {"Host": "localhost"}
    body = {"title": "Vorschau Test", "items": [{"text": "1972-08-18", "category": "geburtsdatum", "date_mode": "verschieben"},
                                                {"text": "Fabian Beispiel", "category": "name", "main": True, "gender": "m"}]}
    up = client.post("/api/upload", data={"file": (io.BytesIO(b"Fabian Beispiel\n1972-08-18"), "x.txt")}, headers=hdr).get_json()
    body["doc_id"] = up["doc_id"]
    pre = client.post("/api/preview", json=body, headers=hdr).get_json()["previews"]
    ex = client.post("/api/export", json=body, headers=hdr).get_json()
    assert pre["Fabian Beispiel"].split()[1].startswith("Muster") and {r["original"]: r["replacement"] for r in ex["mapping"]} == pre, (pre, ex["mapping"])
    print("Teil B (Bewerber/in, Muster-Pseudonyme, Vorschau) geprüft.")


def test_teil_d():
    """Originalansicht: Seitenbilder, Markierungen, Rahmen-Text; Schwärzwerkzeug brennt Stellen fest ein."""
    import fitz
    from anonymisierungstool import server

    src = fitz.open()
    pg = src.new_page()
    pg.insert_text((72, 100), "Anna Beispiel", fontname="helv", fontsize=14)
    pg.insert_text((72, 140), "Geheime Notiz hier", fontname="helv", fontsize=12)
    pg.insert_text((72, 180), "Bleibt stehen", fontname="helv", fontsize=12)
    pg2 = src.new_page()
    pg2.insert_text((72, 100), "Zweite Seite bleibt", fontname="helv", fontsize=12)
    data = src.tobytes()

    client = server.app.test_client()
    client.set_cookie("tok", server.TOKEN, domain="localhost")
    hdr = {"Host": "localhost"}
    up = client.post("/api/upload", data={"file": (io.BytesIO(data), "d.pdf")}, headers=hdr).get_json()
    did = up["doc_id"]
    assert len(up["pages"]) == 2 and abs(up["pages"][0]["w"] - 595) < 5, up["pages"]
    png = client.get("/api/doc/%s/page/0.png" % did, headers=hdr)
    assert png.status_code == 200 and png.data[:4] == b"\x89PNG" and client.get("/api/doc/%s/page/9.png" % did, headers=hdr).status_code == 404
    marks = client.post("/api/doc/%s/marks" % did, json={"terms": ["Anna Beispiel", "Nichtda"]}, headers=hdr).get_json()["marks"]
    assert len(marks) == 1 and marks[0]["term"] == "Anna Beispiel" and marks[0]["page"] == 0
    m = marks[0]
    assert 0.08 < m["x"] < 0.16 and 0.1 < m["y"] < 0.15 and 0.1 < m["w"] < 0.3, m
    # Rahmen um "Geheime Notiz" (Anteile der Seite) liefert den Text darin
    txt = client.post("/api/doc/%s/rect" % did, json={"page": 0, "x0": 0.05, "y0": 0.15, "x1": 0.5, "y1": 0.19}, headers=hdr).get_json()["text"]
    assert txt == "Geheime Notiz hier", txt
    # Klick auf ein einzelnes Wort ("Geheime" liegt bei x = 72 pt = 0.12 der Seitenbreite, y = 140 pt)
    word = client.post("/api/doc/%s/rect" % did, json={"page": 0, "x0": 0.13, "y0": 140 / 842 - 0.005, "x1": 0.13, "y1": 140 / 842 - 0.005, "touch": True}, headers=hdr).get_json()["text"]
    assert word == "Geheime", word
    # Ziehen, das ein Wort nur am Rand streift, nimmt es nicht mit
    part = client.post("/api/doc/%s/rect" % did, json={"page": 0, "x0": 0.05, "y0": 0.15, "x1": 0.16, "y1": 0.19}, headers=hdr).get_json()["text"]
    assert part == "Geheime", part
    # Ohne Markierung kein Text
    assert client.post("/api/doc/%s/rect" % did, json={"page": 0, "x0": 0.7, "y0": 0.7, "x1": 0.9, "y1": 0.9}, headers=hdr).get_json()["text"] == ""

    # Export: ein Fund ersetzt, drei Schwärzungen von Hand (Rechteck, Kreis, Stift) auf Seite 1
    body = {"doc_id": did, "title": "Schwaerzen Test", "items": [{"text": "Anna Beispiel", "category": "name"}], "blackouts": [
        {"page": 0, "shape": "rect", "x0": 0.1, "y0": 0.15, "x1": 0.6, "y1": 0.19},
        {"page": 0, "shape": "ellipse", "x0": 0.1, "y0": 0.2, "x1": 0.4, "y1": 0.25},
        {"page": 0, "shape": "pen", "pts": [[0.1, 0.3], [0.3, 0.31], [0.5, 0.3]], "w": 0.03}]}
    ex = client.post("/api/export", json=body, headers=hdr).get_json()
    assert "error" not in ex and ex["result_pages"] == 2, ex
    out = fitz.open(stream=client.get("/download/%s/doc" % ex["result_id"], headers=hdr).data, filetype="pdf")
    assert len(out) == 2
    assert out[0].get_text().strip() == "" and len(out[0].get_images()) == 1, "Seite 1 muss ein reines Bild sein (kein Text darunter)"
    assert "Zweite Seite bleibt" in out[1].get_text(), "Seite ohne Schwärzung bleibt unverändert"
    pix = out[0].get_pixmap(dpi=72)
    R = out[0].rect
    assert pix.pixel(int(0.3 * R.width), int(0.17 * R.height))[0] < 40, "Rechteck ist nicht schwarz"
    assert pix.pixel(int(0.25 * R.width), int(0.225 * R.height))[0] < 40, "Kreis ist nicht schwarz"
    assert pix.pixel(int(0.3 * R.width), int(0.31 * R.height))[0] < 40, "Stift ist nicht schwarz"
    assert pix.pixel(int(0.8 * R.width), int(0.8 * R.height))[0] > 200, "Rest der Seite muss hell bleiben"
    assert client.get("/api/result/%s/page/1.png" % ex["result_id"], headers=hdr).data[:4] == b"\x89PNG"
    # nur Schwärzungen, kein Fund: erlaubt (z. B. Scan-Seiten); weder Fund noch Schwärzung: Fehler
    only = client.post("/api/export", json={"doc_id": did, "title": "Schwaerzen Test", "items": [], "blackouts": body["blackouts"][:1]}, headers=hdr)
    assert only.status_code == 200
    none = client.post("/api/export", json={"doc_id": did, "title": "Schwaerzen Test", "items": []}, headers=hdr)
    assert none.status_code == 400
    print("Teil D (Originalansicht, Markierungen, Schwärzwerkzeug) geprüft.")


def test_sicherheit():
    """Fremde Anfragen, XML-Entitäten, Metadaten und Escaping."""
    from lxml import etree
    from anonymisierungstool import server
    from anonymisierungstool.documents import _SAFE_XML

    client = server.app.test_client()
    client.set_cookie("tok", server.TOKEN, domain="localhost")
    hdr = {"Host": "localhost:8765"}
    ok = client.post("/api/preview", json={"title": "Sicherheit", "items": []}, headers=dict(hdr, Origin="http://localhost:8765"))
    evil = client.post("/api/preview", json={"title": "Sicherheit", "items": []}, headers=dict(hdr, Origin="http://localhost:9999"))
    assert ok.status_code == 200 and evil.status_code == 403, (ok.status_code, evil.status_code)
    assert client.get("/", headers={"Host": "evil.example"}).status_code == 403          # DNS-Rebinding
    assert client.get("/api/config", headers={"Host": "localhost"}).status_code == 200
    no_cookie = server.app.test_client().get("/api/config", headers={"Host": "localhost"})
    assert no_cookie.status_code == 403

    secret = os.path.join(tempfile.mkdtemp(), "geheim.txt")
    open(secret, "w").write("GEHEIMER-INHALT")
    from pathlib import Path

    xxe = '<?xml version="1.0"?><!DOCTYPE a [<!ENTITY e SYSTEM "%s">]><a>&e;</a>' % Path(secret).as_uri()
    try:
        assert "GEHEIMER-INHALT" not in "".join(etree.fromstring(xxe.encode(), _SAFE_XML).itertext()), "externe Entität wurde gelesen"
    except etree.XMLSyntaxError:
        pass  # Ablehnen ist ebenfalls sicher

    d = docx.Document()
    d.add_paragraph("Text")
    buf = io.BytesIO()
    d.save(buf)
    out, _, _ = DocxDoc(buf.getvalue()).export(None, {}, True)
    import zipfile
    names = zipfile.ZipFile(io.BytesIO(out)).namelist()
    assert not any(n.startswith("docProps/") and ("app" in n or "custom" in n or "thumbnail" in n) for n in names), names
    DocxDoc(out)   # bleibt lesbar
    print("Sicherheitsprüfungen (Origin, Host, XXE, Metadaten) bestanden.")


def test_ablage():
    """Ablage: Übersicht, Dateien eines Titels, Ablageort ändern, sicheres Öffnen."""
    from anonymisierungstool import server, titel

    client = server.app.test_client()
    client.set_cookie("tok", server.TOKEN, domain="localhost")
    hdr = {"Host": "localhost"}
    lib = client.get("/api/library", headers=hdr).get_json()
    assert lib["locked"] is True and lib["out_dir"] == os.environ["ANONYMISIERUNGSTOOL_AUSGABE"]
    mine = [t for t in lib["titles"] if t["name"] == "Schwaerzen Test"]
    assert mine and mine[0]["docs"] == 2 and mine[0]["has_key"] is True and mine[0]["restored"] == 0, lib["titles"]
    files = client.get("/api/title-files", query_string={"title": "Schwaerzen Test"}, headers=hdr).get_json()["files"]
    assert [f["kind"] for f in files] == ["anonymisiert", "anonymisiert", "schluessel"], files
    assert client.post("/api/library/location", json={"path": "/tmp/x"}, headers=hdr).status_code == 400   # per Umgebungsvariable festgelegt

    saved = os.environ.pop("ANONYMISIERUNGSTOOL_AUSGABE")
    try:
        new_dir = tempfile.mkdtemp()
        moved = client.post("/api/library/location", json={"path": new_dir}, headers=hdr).get_json()
        assert moved["out_dir"] == new_dir and moved["titles"] == [] and moved["locked"] is False
        assert client.post("/api/library/location", json={"path": "relativ/pfad"}, headers=hdr).status_code == 400
        assert client.post("/api/library/location", json={"path": "/proc/nicht/schreibbar"}, headers=hdr).status_code == 400
        up = client.post("/api/upload", data={"file": (io.BytesIO(b"Anna Beispiel arbeitet hier."), "a.txt")}, headers=hdr).get_json()
        ex = client.post("/api/export", json={"doc_id": up["doc_id"], "title": "Neuer Ort", "items": [{"text": "Anna Beispiel", "category": "name"}]}, headers=hdr).get_json()
        assert ex["folder"] == os.path.join(new_dir, "Neuer Ort"), ex["folder"]
        assert client.post("/api/open-file", json={"title": "Neuer Ort", "file": "../x"}, headers=hdr).status_code == 400
        assert client.post("/api/open-file", json={"title": "Neuer Ort", "file": "gibtsnicht.docx"}, headers=hdr).status_code == 400
        reset = client.post("/api/library/location", json={"path": ""}, headers=hdr).get_json()
        assert reset["out_dir"] == str(titel.default_dir())
    finally:
        os.environ["ANONYMISIERUNGSTOOL_AUSGABE"] = saved
        from anonymisierungstool import store
        store.save_settings({})
    print("Ablage (Übersicht, Dateien, Ablageort ändern) geprüft.")


def test_originale():
    """Original wird beim Laden im Titel-Ordner abgelegt, der Stand automatisch gesichert und lässt sich weiter bearbeiten."""
    from anonymisierungstool import server, titel

    client = server.app.test_client()
    client.set_cookie("tok", server.TOKEN, domain="localhost")
    hdr = {"Host": "localhost"}
    data = b"Anna Beispiel schreibt an Kowalski Bau GmbH."
    up = client.post("/api/upload", data={"file": (io.BytesIO(data), "../../Brief Anna.txt"), "title": "Original Test"}, headers=hdr).get_json()
    assert up["original"] == "Brief Anna.txt", up["original"]            # keine Pfadanteile
    d = titel.folder("Original Test") / "Originale"
    assert (d / "Brief Anna.txt").read_bytes() == data
    again = client.post("/api/upload", data={"file": (io.BytesIO(data), "Brief Anna.txt"), "title": "Original Test"}, headers=hdr).get_json()
    other = client.post("/api/upload", data={"file": (io.BytesIO(data + b" Neu."), "Brief Anna.txt"), "title": "Original Test"}, headers=hdr).get_json()
    assert again["original"] == "Brief Anna.txt" and other["original"] == "Brief Anna_2.txt"
    # ohne Titel wird nichts gespeichert (nur für Tests/Altfälle)
    assert client.post("/api/upload", data={"file": (io.BytesIO(data), "x.txt")}, headers=hdr).get_json()["original"] is None
    assert client.post("/api/upload", data={"file": (io.BytesIO(data), "x.txt"), "title": "  "}, headers=hdr).get_json()["original"] is None  # leerer Titel = kein Titel

    state = {"findings": [{"id": 0, "text": "Anna Beispiel", "category": "name", "checked": True, "action": "pseudonym", "override": "Kandidatin", "date_mode": "jahr"}],
             "blackouts": [{"page": 0, "shape": "rect", "x0": 0.1, "y0": 0.1, "x1": 0.2, "y1": 0.2}], "rm": False, "done": False, "unerwuenscht": "x"}
    assert client.post("/api/doc/%s/state" % up["doc_id"], json=state, headers=hdr).get_json()["saved"] is True
    assert client.post("/api/doc/nichtda/state", json=state, headers=hdr).get_json()["saved"] is False
    lib = {t["name"]: t for t in client.get("/api/library", headers=hdr).get_json()["titles"]}["Original Test"]
    assert lib["originals"] == 2 and lib["open"] == 2, lib
    info = client.get("/api/title-files", query_string={"title": "Original Test"}, headers=hdr).get_json()
    assert [o["name"] for o in info["originals"]] == ["Brief Anna.txt", "Brief Anna_2.txt"] and info["originals"][0]["has_state"] is True

    back = client.post("/api/resume", json={"title": "Original Test", "original": "Brief Anna.txt"}, headers=hdr).get_json()
    assert back["state"]["findings"][0]["override"] == "Kandidatin" and back["state"]["rm"] is False and "unerwuenscht" not in back["state"], back["state"]
    assert back["text"].startswith("Anna Beispiel") and back["original"] == "Brief Anna.txt"
    assert client.post("/api/resume", json={"title": "Original Test", "original": "../x"}, headers=hdr).status_code == 400
    state["done"] = True
    client.post("/api/doc/%s/state" % back["doc_id"], json=state, headers=hdr)
    lib = {t["name"]: t for t in client.get("/api/library", headers=hdr).get_json()["titles"]}["Original Test"]
    assert lib["open"] == 1, lib
    assert client.post("/api/open-file", json={"title": "Original Test", "file": "../Brief Anna.txt", "sub": "Originale"}, headers=hdr).status_code == 400
    t = client.get("/api/original-thumb?title=Original+Test&file=Brief+Anna.txt", headers=hdr)
    assert t.status_code == 200 and t.data[:4] == b"\x89PNG", t.status_code
    assert client.get("/api/original-thumb?title=Original+Test&file=..%2Fx", headers=hdr).status_code == 404
    assert client.get("/api/original-thumb?sub=ergebnis&title=Original+Test&file=..%2Fx", headers=hdr).status_code == 404
    print("Originale und Bearbeitungsstand (speichern, weiter bearbeiten, Vorschaubild) geprüft.")


def test_namenszeile_gemischt():
    from anonymisierungstool.detect import analyze_all

    t = "ANNA oBEISPIEL-MUSTERLE\nMusterweg 5, 80331 München | anna.obeispiel@example.org\nBanking School-September 2018"
    f, _ = analyze_all(t)
    names = {x["text"]: x for x in f if x["category"] == "name"}
    assert "ANNA oBEISPIEL-MUSTERLE" in names and names["ANNA oBEISPIEL-MUSTERLE"]["checked"], names
    assert not any("Banking" in x or "September" in x for x in names), names
    print("Namenszeile in gemischter Schreibweise, Fehlfund Institution geprüft.")


def test_schreibweisen():
    """Name in GROSSBUCHSTABEN oder klein: eigener Fund, Pseudonym übernimmt die Schreibweise, Rückspielen klappt."""
    import fitz

    from anonymisierungstool.detect import analyze_all
    from anonymisierungstool.restore import build_reverse, regex_for
    from anonymisierungstool.textutil import replace_in_text

    text = "FRITZ MÜLLER\nHerr Fritz Müller wohnt in Berlin. Kontakt: fritz.mueller@example.org\nmüller schreibt, Bewerbung von FRITZ."
    found, _ = analyze_all(text)
    names = {f["text"] for f in found if f["category"] == "name"}
    assert {"Fritz Müller", "FRITZ MÜLLER", "MÜLLER", "FRITZ"} <= names, names
    base = {f["text"]: f for f in found}
    assert base["FRITZ MÜLLER"]["checked"] == base["Fritz Müller"]["checked"]
    items = [Item(f["text"], f["category"], gender=f.get("gender", "")) for f in found if f["category"] == "name"]
    key = {"entries": []}
    m, _e = build_mapping(items, key=key)
    v, n = m["Fritz Müller"].split()
    assert m["FRITZ MÜLLER"] == "%s %s" % (v.upper(), n.upper()) and m["MÜLLER"] == n.upper() and m["FRITZ"] == v.upper(), m
    kl, _e = build_mapping([Item("Fritz Müller", "name"), Item("müller", "name")], key={"entries": []})
    assert kl["müller"] == kl["Fritz Müller"].split()[1].lower(), kl

    d = docx.Document()
    d.add_paragraph("FRITZ MÜLLER")
    d.add_paragraph("Herr Fritz Müller, bekannt als müller.")
    buf = io.BytesIO()
    d.save(buf)
    wdoc = open_document("x.docx", buf.getvalue())
    wf, _ = analyze_all(wdoc.text(), wdoc.alt_texts())
    wi = [Item(f["text"], f["category"]) for f in wf if f["category"] == "name"]
    wk = {"entries": []}
    wm, _e = build_mapping(wi, key=wk)
    out, ext, _ = wdoc.export(term_regex(wm), wm, True)
    assert not wdoc.leftover_check(out, ext, list(wm))
    rtxt = open_document("y.docx", out).text()
    assert "MÜLLER" not in rtxt and "müller" not in rtxt and "MUSTER" in rtxt.upper(), rtxt
    rev, _lost = build_reverse(wk)
    back = replace_in_text(rtxt, regex_for(rev), rev)
    assert "FRITZ MÜLLER" in back and "müller" in back and "Fritz Müller" in back, back

    pd = fitz.open()
    pg = pd.new_page()
    pg.insert_text((50, 60), "FRITZ MUELLER", fontname="hebo", fontsize=16)
    pg.insert_text((50, 100), "Herr Fritz Mueller wohnt hier.", fontname="helv", fontsize=11)
    pdoc = open_document("x.pdf", pd.tobytes())
    pf, _ = analyze_all(pdoc.text(), pdoc.alt_texts())
    assert {"FRITZ MUELLER", "Fritz Mueller"} <= {f["text"] for f in pf if f["category"] == "name"}, [f["text"] for f in pf]
    pm, _e = build_mapping([Item(f["text"], f["category"]) for f in pf if f["category"] == "name"], key={"entries": []})
    pout, pext, _ = pdoc.export(term_regex(pm), pm, True)
    ptxt = fitz.open(stream=pout, filetype="pdf")[0].get_text()
    assert "MUELLER" not in ptxt and "Mueller" not in ptxt and "FRITZ" not in ptxt and "MUSTER" in ptxt, ptxt
    from anonymisierungstool import server
    c = server.app.test_client(); c.set_cookie("tok", server.TOKEN, domain="localhost"); hdr = {"Host": "localhost"}
    up = c.post("/api/upload", data={"file": (io.BytesIO(sample_docx()), "cv.docx")}, headers=hdr).get_json()
    h = c.get("/api/doc/%s/html" % up["doc_id"], headers=hdr).get_json()["html"]
    assert 'class="wpage"' in h and "<table" in h and "Anna" in h and "Müller &amp; Söhne" in h, h[:300]
    assert c.get("/api/doc/%s/html" % up["doc_id"].replace("a", "b"), headers=hdr).status_code in (200, 404)
    ex = c.post("/api/export", json={"doc_id": up["doc_id"], "title": "Wordansicht", "items": [{"text": "Anna Beispiel", "category": "name"}]}, headers=hdr).get_json()
    rh = c.get("/api/result/%s/html" % ex["result_id"], headers=hdr).get_json()["html"]
    assert "Anna Beispiel" not in rh and "Muster" in rh
    print("Schreibweisen (Versalien/klein, Word, PDF, Rückspielen) und Word-Ansicht geprüft.")


def test_erkennung_luecken():
    """Lücken aus dem KI-Lesetest: Name in der Fußzeile, Zeiträume sind keine Telefonnummern, "Str." mit Großbuchstaben."""
    from anonymisierungstool.detect import analyze_all

    t = "\n".join(["Berufserfahrung"] + ["Tätigkeit %d" % i for i in range(14)] +
                   ["Musterstraße 1, 12345 Musterstadt\tTel. 0175 62 90 168\tmuster@example.com", "ANNA BEISPIELFRAU", "Seite 2"])
    f, _ = analyze_all(t)
    names = {x["text"]: x for x in f if x["category"] == "name"}
    assert "ANNA BEISPIELFRAU" in names and names["ANNA BEISPIELFRAU"]["checked"], "Name in der Fußzeile neben Kontaktdaten"
    f, _ = analyze_all("09/77 - 07/85 Schule\n06/23-09/24 Praktikum\nRegensburger Str. 12, Nürnberg\nTel. 0175 62 90 168")
    cats = [(x["category"], x["text"]) for x in f]
    assert ("telefon", "0175 62 90 168") in cats and not any(c == "telefon" and "/" in x for c, x in cats), cats
    assert ("adresse", "Regensburger Str. 12") in cats, cats
    f, _ = analyze_all("Regensburger Str. 17990478, Nürnberg\nIch lebe gern in Nürnberg.")
    c2 = {x["text"]: x for x in f}
    assert c2["Nürnberg"]["category"] == "adresse" and c2["Nürnberg"]["checked"] and c2["Nürnberg"]["count"] == 2, "Ort hinter der Straße gehört zur Adresse"
    f, _ = analyze_all("Einsatz von Bitrix und HIV-Programme bei Scrum", [])
    assert not [x for x in f if x["category"] in ("firma", "ort") and x["text"] in ("Bitrix", "HIV", "Scrum")]
    print("Erkennungslücken (Fußzeilen-Name, Zeiträume, Str.) geprüft.")


def test_neu_gestalteter_lebenslauf():
    """Zweite, neu gestaltete Fassung im selben Titel: gleiche Pseudonyme, Struktur bleibt, Rückspielen klappt."""
    import re
    import zipfile

    from anonymisierungstool import server, titel

    c = server.app.test_client(); c.set_cookie("tok", server.TOKEN, domain="localhost"); hdr = {"Host": "localhost"}

    def v1():
        d = docx.Document()
        t = d.add_table(rows=1, cols=2)
        left, right = t.rows[0].cells
        left.paragraphs[0].add_run("Anna Beispiel").bold = True
        left.add_paragraph("Musterweg 12, 80331 München")
        left.add_paragraph("anna.beispiel@example.org")
        right.paragraphs[0].add_run("03/2019 – 08/2021 Projektleiterin bei Kowalski Bau GmbH, Warschau")
        b = io.BytesIO(); d.save(b); return b.getvalue()

    def v2():
        d = docx.Document()
        d.add_heading("ANNA BEISPIEL", 0)
        d.add_paragraph("Kontakt: anna.beispiel@example.org · Lindenallee 5, 86150 Augsburg")
        d.add_paragraph("03/2019 – 08/2021 Projektleiterin bei Kowalski Bau GmbH, Warschau · Seit 2021 Nordlicht Consulting AG, Hamburg")
        b = io.BytesIO(); d.save(b); return b.getvalue()

    def skel(b):
        return re.sub(r"(<w:t[^>]*>)[^<]*", r"\1", zipfile.ZipFile(io.BytesIO(b)).read("word/document.xml").decode()).replace(' xml:space="preserve"', "")

    maps, outs = [], []
    for name, data in (("cv_alt.docx", v1()), ("cv_neu.docx", v2())):
        up = c.post("/api/upload", data={"file": (io.BytesIO(data), name), "title": "Neu gestaltet"}, headers=hdr).get_json()
        assert not [f for f in up["findings"] if f["text"] in ("anna", "beispiel")], "Teile von E-Mail-Adressen sind keine eigenen Funde"
        items = [{"text": f["text"], "category": f["category"], "gender": f.get("gender", "")} for f in up["findings"] if f["checked"] or f["category"] in ("firma", "ort")]
        ex = c.post("/api/export", json={"doc_id": up["doc_id"], "title": "Neu gestaltet", "items": items}, headers=hdr).get_json()
        assert not ex["leftover"] and ex["file"].startswith(name[:-5]), ex
        out = (titel.folder("Neu gestaltet") / ex["file"]).read_bytes()
        assert skel(out) == skel(data), "Struktur der Word-Datei hat sich verändert"
        maps.append({m["original"]: m["replacement"] for m in ex["mapping"]}); outs.append(out)
    assert maps[0]["Kowalski Bau GmbH"] == maps[1]["Kowalski Bau GmbH"] == "Musterfirma A" and maps[1]["Nordlicht Consulting AG"] == "Musterfirma B"
    assert maps[1]["ANNA BEISPIEL"] == maps[0]["Anna Beispiel"].upper(), maps
    back = c.post("/api/restore", data={"file": (io.BytesIO(outs[1]), "ki.docx"), "title": "Neu gestaltet"}, headers=hdr).get_json()
    txt = DocxDoc((titel.folder("Neu gestaltet") / back["file"]).read_bytes()).text()
    assert back["kind"] == "docx" and "Anna" not in c.get("/api/result/%s/html" % back["result_id"], headers=hdr).get_json()["html"].replace("ANNA BEISPIEL", "")
    assert "ANNA BEISPIEL" in txt and "Nordlicht Consulting AG" in txt and "Lindenallee 5" in txt and not back["unknown"], (txt, back)
    print("Neu gestalteter Lebenslauf im selben Titel (Pseudonyme, Struktur, Rückspielen) geprüft.")


def test_wordansicht_elemente():
    from anonymisierungstool.docx_html import render

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Kopfzeile Test"
    d.add_paragraph("Liste A", style="List Bullet")
    d.add_paragraph("Nummer 1", style="List Number")
    d.add_paragraph("Nummer 2", style="List Number")
    d.add_paragraph().add_run("fett").bold = True
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(0, 0).text = "Kopf"
    buf = io.BytesIO()
    d.save(buf)
    h = render(DocxDoc(buf.getvalue()))
    for part in ("Kopfzeile Test", "•", "1.", "2.", "font-weight:700", 'colspan="2"', "border-top:1px solid"):
        assert part in h, part
    print("Word-Ansicht (Kopfzeile, Listen, Tabellen, Schrift) geprüft.")


def test_live_vorschau():
    """Vorschau mit aktuellem Stand (Word und PDF) und mehrfaches Erstellen mit anderer Auswahl."""
    from anonymisierungstool import server

    c = server.app.test_client(); c.set_cookie("tok", server.TOKEN, domain="localhost"); hdr = {"Host": "localhost"}
    items = [{"text": "Anna Beispiel", "category": "name"}, {"text": "Kowalski Bau GmbH", "category": "firma"}]
    up = c.post("/api/upload", data={"file": (io.BytesIO(sample_docx()), "cv.docx")}, headers=hdr).get_json()
    body = {"doc_id": up["doc_id"], "title": "Live Vorschau", "items": items}
    p1 = c.post("/api/doc/%s/preview-render" % up["doc_id"], json=body, headers=hdr).get_json()
    assert "Anna Beispiel" not in p1["html"] and p1["previews"]["Kowalski Bau GmbH"].startswith("Musterfirma"), p1["previews"]
    body["items"] = items[:1]   # Firma wieder abgehakt: Original muss zurückkommen (frische Kopie, nicht das schon veränderte Dokument)
    p2 = c.post("/api/doc/%s/preview-render" % up["doc_id"], json=dict(body, prev_rid=p1["rid"]), headers=hdr).get_json()
    assert "Kowalski Bau GmbH" in p2["html"] and "Anna Beispiel" not in p2["html"]
    assert c.get("/api/result/%s/html" % p1["rid"], headers=hdr).status_code == 404, "alte Vorschau wird aufgeräumt"
    upd = c.post("/api/upload", data={"file": (io.BytesIO(sample_pdf()), "cv.pdf")}, headers=hdr).get_json()
    pp = c.post("/api/doc/%s/preview-render" % upd["doc_id"], json={"title": "Live Vorschau", "items": [{"text": "Anna Beispiel", "category": "name"}], "remove_images": True}, headers=hdr).get_json()
    assert pp["pages"] == 1 and pp["previews"]["Anna Beispiel"]
    marks = c.post("/api/result/%s/marks" % pp["rid"], json={"terms": [pp["previews"]["Anna Beispiel"]]}, headers=hdr).get_json()["marks"]
    assert marks and c.get("/api/result/%s/page/0.png" % pp["rid"], headers=hdr).data[:4] == b"\x89PNG"
    rr = c.post("/api/result/%s/rect" % pp["rid"], json={"page": 0, "x0": 0, "y0": 0, "x1": 1, "y1": 1}, headers=hdr).get_json()["text"]
    assert pp["previews"]["Anna Beispiel"].split()[0] in rr and "Beispiel" not in rr, rr
    print("Live-Vorschau (Word, PDF, Markierungen, mehrfach) geprüft.")


def test_namen_ueberall():
    """Jeder Teil eines gewählten Namens wird in jeder Schreibweise ersetzt, auch wenn die Erkennung ihn anders einordnete."""
    from anonymisierungstool import server, titel
    from anonymisierungstool.detect import _add_case_variants

    text = "Amir Degan\nAMIR\nDEGAN\nHerr degan und Amir\nMail: amir.degan@example.org"
    base = {"category": "name", "source": "modell", "confident": True, "checked": True, "count": 1, "context": "", "main": False}
    finds = [dict(base, text="Amir Degan"), dict(base, text="AMIR", category="firma", checked=False, confident=False)]
    _add_case_variants(finds, text)
    cat = {f["text"]: f["category"] for f in finds}
    assert cat["AMIR"] == "name" and cat["DEGAN"] == "name" and cat["degan"] == "name" and cat["Amir"] == "name" and "amir" not in cat, cat
    c = server.app.test_client(); c.set_cookie("tok", server.TOKEN, domain="localhost"); hdr = {"Host": "localhost"}
    up = c.post("/api/upload", data={"file": (io.BytesIO(text.encode()), "x.txt"), "title": "Namen überall"}, headers=hdr).get_json()
    ex = c.post("/api/export", json={"doc_id": up["doc_id"], "title": "Namen überall", "items": [{"text": "Amir Degan", "category": "name"}, {"text": "amir.degan@example.org", "category": "email"}]}, headers=hdr).get_json()
    out = (titel.folder("Namen überall") / ex["file"]).read_text(encoding="utf-8")
    assert not re.search(r"amir|degan", out, re.I), out
    m = {r["original"]: r["replacement"] for r in ex["mapping"]}
    assert m["AMIR"] == m["Amir"].upper() and m["DEGAN"] == m["Degan"].upper() if "Degan" in m else m["AMIR"] == m["Amir"].upper(), m
    print("Namensteile in jeder Schreibweise (auch vom Modell anders eingeordnet) geprüft.")


def test_loeschen():
    """Löschen von Titel, Original und erstellter Datei; Schlüssel und fremde Pfade sind geschützt."""
    from anonymisierungstool import server, titel

    c = server.app.test_client(); c.set_cookie("tok", server.TOKEN, domain="localhost"); hdr = {"Host": "localhost"}
    up = c.post("/api/upload", data={"file": (io.BytesIO(b"Anna Beispiel wohnt hier."), "brief.txt"), "title": "Weg damit"}, headers=hdr).get_json()
    ex = c.post("/api/export", json={"doc_id": up["doc_id"], "title": "Weg damit", "items": [{"text": "Anna Beispiel", "category": "name"}]}, headers=hdr).get_json()
    d = titel.folder("Weg damit")
    assert (d / "Originale" / "brief.txt").is_file() and (d / ex["file"]).is_file()
    post = lambda body: c.post("/api/delete", json=body, headers=hdr)
    assert post({"title": "Weg damit", "kind": "result", "file": "SCHLUESSEL_nicht_in_die_KI_laden.anonkey"}).status_code == 400, "Schlüssel nicht einzeln löschbar"
    assert post({"title": "Weg damit", "kind": "result", "file": "../x.txt"}).status_code == 400
    assert post({"title": "Weg damit", "kind": "result", "file": ex["file"]}).status_code == 200 and not (d / ex["file"]).exists()
    assert post({"title": "Weg damit", "kind": "original", "file": "brief.txt"}).status_code == 200 and not (d / "Originale" / "brief.txt").exists()
    assert post({"title": "Weg damit", "kind": "original", "file": "brief.txt"}).status_code == 400
    assert c.post("/api/doc/%s/state" % up["doc_id"], json={"findings": []}, headers=hdr).get_json().get("saved") is False, "kein Wiederanlegen nach dem Löschen"
    assert post({"title": "Weg damit", "kind": "title"}).status_code == 200 and not d.exists()
    assert post({"title": "Weg damit", "kind": "title"}).status_code == 400
    assert post({"title": "..", "kind": "title"}).status_code == 400 and titel.base_dir().is_dir()
    print("Löschen (Titel, Original, Datei, Schutz) geprüft.")


def main():
    data = sample_docx()
    doc = open_document("cv.docx", data)
    text = doc.text()
    print("--- TEXT ---\n" + text)
    findings, warn = analyze(text)
    print("--- FUNDE (warn=%s) ---" % warn)
    for f in findings:
        print("%-14s %-34s %s %s%s" % (f["category"], f["text"], "[x]" if f["checked"] else "[ ]", f["source"], "" if f["confident"] else " (unsicher)"))

    items = [Item(f["text"], f["category"]) for f in findings if f["checked"]]
    items.append(Item("Kowalski Bau GmbH", "firma"))
    mapping, entries = build_mapping(items, "verschieben")
    print("--- MAPPING ---")
    for k, v in mapping.items():
        print("%-34s -> %s" % (k, v))

    out, ext, _ = DocxDoc(data).export(term_regex(mapping), mapping, True)
    res = DocxDoc(out)
    print("--- AUSGABE (docx) ---\n" + res.text())
    left = DocxDoc(out).leftover_check(out, "docx", list(mapping))
    print("Übrig:", left)
    assert "Anna" not in res.all_text() and "Beispiel" not in res.all_text(), "Name noch enthalten"
    import zipfile as _z
    assert any(n.endswith(".png") for n in _z.ZipFile(io.BytesIO(out)).namelist() if "media" in n) or res.image_count == 0, "Platzhalterbild fehlt"
    assert not left
    assert res.doc.core_properties.author == ""
    test_pdf(mapping)
    test_scan_pages()
    test_teil_a()
    test_teil_c()
    test_teil_b()
    test_namenszeile_gemischt()
    test_namen_ueberall()
    test_schreibweisen()
    test_erkennung_luecken()
    test_neu_gestalteter_lebenslauf()
    test_wordansicht_elemente()
    test_live_vorschau()
    test_loeschen()
    test_teil_d()
    test_sicherheit()
    test_ablage()
    test_originale()

    assert load_key(dump_key({"entries": entries}))["entries"] == entries
    try:
        load_key(b"ANK1" + b"x" * 40)
        raise AssertionError("alte verschlüsselte Schlüsseldatei wurde akzeptiert")
    except ValueError:
        pass
    test_titel()
    print("\nAlle Prüfungen bestanden.")


if __name__ == "__main__":
    main()
